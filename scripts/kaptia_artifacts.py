"""Export Kaptia content into validated, readable deliverable formats."""
import csv
import hashlib
import io
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "artifacts"
SUPPORTED = {"md", "txt", "json", "csv", "docx", "pdf", "pptx"}
MAX_TEXT_CHARS = 1_000_000


def _safe_name(name):
    if not isinstance(name, str) or not name.strip() or len(name) > 100:
        raise ValueError("檔名必須為 1–100 字元")
    name = re.sub(r"[^A-Za-z0-9_\-\u3400-\u9fff]+", "-", name.strip()).strip("-")
    if not name or name in {".", ".."}:
        raise ValueError("檔名無效")
    return name


def export_artifact(content, output_type="md", filename="kaptia-output", data=None):
    if not isinstance(content, str) or len(content) > MAX_TEXT_CHARS:
        raise ValueError("content 必須是字串且不超過 1,000,000 字元")
    if output_type not in SUPPORTED:
        raise ValueError(f"不支援的格式；可用格式：{sorted(SUPPORTED)}")
    if output_type in {"json", "csv"} and data is None:
        raise ValueError(f"{output_type} 格式需要提供結構化 data")
    base = _safe_name(filename)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    target = (OUTPUT_DIR / f"{base}.{output_type}").resolve()
    if OUTPUT_DIR.resolve() not in target.parents:
        raise ValueError("輸出路徑無效")
    if output_type in {"md", "txt"}:
        payload = content.encode("utf-8")
    elif output_type == "json":
        payload = (json.dumps(data, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    elif output_type == "csv":
        if not isinstance(data, list) or any(not isinstance(row, dict) for row in data):
            raise ValueError("CSV data 必須是 dict 清單")
        fields = list(dict.fromkeys(key for row in data for key in row.keys()))
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=fields, extrasaction="raise")
        if fields:
            writer.writeheader()
            writer.writerows(data)
        payload = buffer.getvalue().encode("utf-8-sig")
    elif output_type == "docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX 匯出需要安裝 python-docx") from exc
        document = Document()
        for block in re.split(r"\n\s*\n", content.strip()):
            lines = block.splitlines()
            if lines and lines[0].startswith("#"):
                level = min(3, max(1, len(lines[0]) - len(lines[0].lstrip("#"))))
                document.add_heading(lines[0].lstrip("# ").strip(), level=level)
                if len(lines) > 1:
                    document.add_paragraph("\n".join(lines[1:]))
            elif block.strip():
                document.add_paragraph(block)
        stream = io.BytesIO()
        document.save(stream)
        payload = stream.getvalue()
    elif output_type == "pdf":
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.units import mm
        except ImportError as exc:
            raise RuntimeError("PDF 匯出需要安裝 reportlab") from exc
        # Built-in Helvetica cannot render all CJK glyphs. Fail closed instead of
        # silently producing a PDF with missing Traditional Chinese characters.
        if re.search(r"[\u3400-\u9fff]", content):
            raise RuntimeError("內容包含中文；請先配置並註冊支援繁體中文的 TTF 字型，避免 PDF 缺字")
        stream = io.BytesIO()
        styles = getSampleStyleSheet()
        story = []
        for line in content.splitlines():
            escaped = (line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") or " ")
            story.append(Paragraph(escaped, styles["BodyText"]))
            story.append(Spacer(1, 2 * mm))
        SimpleDocTemplate(stream, pagesize=A4).build(story)
        payload = stream.getvalue()
    elif output_type == "pptx":
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError as exc:
            raise RuntimeError("PPTX 匯出需要安裝 python-pptx") from exc
        presentation = Presentation()
        for block in re.split(r"\n\s*\n", content.strip()):
            if not block.strip():
                continue
            lines = block.splitlines()
            layout = presentation.slide_layouts[1 if len(presentation.slide_layouts) > 1 else 0]
            slide = presentation.slides.add_slide(layout)
            slide.shapes.title.text = lines[0].lstrip("# ").strip()[:200]
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = "\n".join(lines[1:])[:4000]
        stream = io.BytesIO()
        presentation.save(stream)
        payload = stream.getvalue()
    target.write_bytes(payload)
    reread = target.read_bytes()
    if reread != payload:
        target.unlink(missing_ok=True)
        raise IOError("交付檔案讀回驗證失敗")
    if output_type == "json":
        json.loads(reread.decode("utf-8"))
    elif output_type == "csv":
        list(csv.reader(io.StringIO(reread.decode("utf-8-sig"))))
    elif output_type == "docx":
        from docx import Document
        Document(io.BytesIO(reread))
    elif output_type == "pptx":
        from pptx import Presentation
        Presentation(io.BytesIO(reread))
    elif output_type == "pdf":
        if not reread.startswith(b"%PDF-"):
            raise IOError("輸出不是有效 PDF")
    return {
        "ok": True, "path": str(target), "format": output_type,
        "bytes": len(payload), "sha256": hashlib.sha256(payload).hexdigest(),
        "readback_verified": True,
    }
