import argparse
import json
from pathlib import Path

def inspect_pptx(path):
    from pptx import Presentation
    prs = Presentation(str(path))
    slides = []
    for number, slide in enumerate(prs.slides, 1):
        items = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                items.append({"type": "text", "text": shape.text})
            if getattr(shape, "has_table", False):
                items.append({"type": "table", "rows": [[c.text for c in row.cells] for row in shape.table.rows]})
            if getattr(shape, "has_chart", False):
                items.append({"type": "chart", "name": shape.name})
        slides.append({"slide": number, "items": items})
    return {"file": str(path), "slide_count": len(prs.slides), "slides": slides}

def replace_pptx_text(source, output, old, new):
    from pptx import Presentation
    prs = Presentation(str(source))
    count = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        count += run.text.count(old)
                        run.text = run.text.replace(old, new)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        count += cell.text.count(old)
                        cell.text = cell.text.replace(old, new)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return {"output": str(output), "replacements": count, "slides": len(prs.slides)}

def create_pptx(output, title, bullets):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "由 Kaptia 建立，請檢查內容"
    for offset in range(0, len(bullets), 5):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "重點 " + str(offset // 5 + 1)
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for idx, bullet in enumerate(bullets[offset:offset + 5]):
            para = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
            para.text = bullet
            para.font.size = Pt(24)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return {"output": str(output), "slides": len(prs.slides)}

def inspect_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(str(path), strict=True)
    return {"file": str(path), "page_count": len(reader.pages), "encrypted": reader.is_encrypted,
            "pages": [{"page": i, "text": (page.extract_text() or "").strip()} for i, page in enumerate(reader.pages, 1)],
            "note": "掃描影像頁可能沒有可抽取文字；目前未提供 OCR。"}

def merge_pdf(inputs, output):
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    for path in inputs:
        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted:
            raise ValueError("PDF 已加密，停止處理：" + str(path))
        for page in reader.pages:
            writer.add_page(page)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    with open(output, "wb") as stream:
        writer.write(stream)
    return {"output": str(output), "pages": len(writer.pages), "inputs": len(inputs)}

def split_pdf(path, output_dir):
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader(str(path), strict=True)
    if reader.is_encrypted:
        raise ValueError("PDF 已加密，停止處理。")
    folder = Path(output_dir)
    folder.mkdir(parents=True, exist_ok=True)
    outputs = []
    for index, page in enumerate(reader.pages, 1):
        writer = PdfWriter()
        writer.add_page(page)
        target = folder / (Path(path).stem + "-page-" + str(index).zfill(3) + ".pdf")
        with target.open("wb") as stream:
            writer.write(stream)
        outputs.append(str(target))
    return {"output_dir": str(folder), "page_count": len(outputs), "files": outputs}

def main():
    parser = argparse.ArgumentParser(description="Kaptia 本機 PPTX/PDF 工具")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("inspect-pptx"); p.add_argument("input")
    p = sub.add_parser("replace-pptx-text"); p.add_argument("input"); p.add_argument("output"); p.add_argument("--old", required=True); p.add_argument("--new", required=True)
    p = sub.add_parser("create-pptx"); p.add_argument("output"); p.add_argument("--title", required=True); p.add_argument("--bullets", nargs="*", default=[])
    p = sub.add_parser("inspect-pdf"); p.add_argument("input")
    p = sub.add_parser("merge-pdf"); p.add_argument("output"); p.add_argument("inputs", nargs="+")
    p = sub.add_parser("split-pdf"); p.add_argument("input"); p.add_argument("output_dir")
    args = parser.parse_args()
    if args.command == "inspect-pptx": result = inspect_pptx(args.input)
    elif args.command == "replace-pptx-text": result = replace_pptx_text(args.input, args.output, args.old, args.new)
    elif args.command == "create-pptx": result = create_pptx(args.output, args.title, args.bullets)
    elif args.command == "inspect-pdf": result = inspect_pdf(args.input)
    elif args.command == "merge-pdf": result = merge_pdf(args.inputs, args.output)
    else: result = split_pdf(args.input, args.output_dir)
    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
